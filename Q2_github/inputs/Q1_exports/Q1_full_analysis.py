"""FMAT3888 Question 1(a)-(e): reproducible monthly parameter estimation.

Run: python Q1_full_analysis.py [--input PATH] [--output-dir PATH]
Dependencies: pandas, numpy, matplotlib, openpyxl (Excel I/O).
No price filling, zero-return deletion, covariance repair or optimisation.
Results are static snapshots; rerun this script to refresh the workbook.
"""
from pathlib import Path
import argparse
import hashlib
import platform
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.ticker import PercentFormatter
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

FREQUENCIES = ("Daily", "Weekly", "Monthly")


# Q1(a): parse source sheets and audit observations
def read_prices(path, sheet):
    raw = pd.read_excel(path, sheet_name=sheet, header=None)
    names = raw.iloc[3, 1:]
    if len(names) != 9 or names.isna().any():
        raise ValueError(f"{sheet}: expected nine named assets in Excel row 4")
    names = names.astype(str).str.strip().tolist()
    if len(set(names)) != 9 or any(n.startswith("Unnamed") for n in names):
        raise ValueError(f"{sheet}: invalid or duplicate asset names")
    if raw.iloc[5, 0] != "Dates" or not (raw.iloc[5, 1:] == "PX_LAST").all():
        raise ValueError(f"{sheet}: unexpected Bloomberg header structure")
    body = raw.iloc[6:].dropna(how="all").copy()
    dates = pd.to_datetime(body.iloc[:, 0], errors="coerce")
    if dates.isna().any() or dates.duplicated().any():
        raise ValueError(f"{sheet}: invalid or duplicate observation dates")
    values = body.iloc[:, 1:].copy()
    numeric = values.apply(pd.to_numeric, errors="coerce")
    # Excel error cells are read as NaN. Only recognise explicit Bloomberg NA text.
    known_na = values.isna() | values.astype(str).apply(
        lambda s: s.str.strip().isin(["#N/A N/A", "#N/A", "N/A", ""]))
    if (numeric.isna() & ~known_na).any().any():
        raise ValueError(f"{sheet}: unrecognised nonnumeric price token")
    numeric.columns = names
    numeric.index = pd.DatetimeIndex(dates, name="Date")
    numeric = numeric.sort_index()
    arr = numeric.to_numpy()
    if np.isinf(arr).any() or ((arr <= 0) & np.isfinite(arr)).any():
        raise ValueError(f"{sheet}: nonpositive or infinite index levels")
    return numeric


def audit_prices(data):
    records = []
    for frequency, prices in data.items():
        for asset, s in prices.items():
            valid = s.dropna()
            if valid.empty:
                raise ValueError(f"{frequency}/{asset}: no valid prices")
            first, last = valid.index[0], valid.index[-1]
            leading = int((s.index < first).sum())
            trailing = int((s.index > last).sum())
            internal = int(s.loc[first:last].isna().sum())
            # Successive valid observations, matching the preliminary audit.
            # This does not treat missing observations as new information.
            changes = valid.diff().iloc[1:].ne(0)
            changed_dates = valid.index[1:][changes.to_numpy()]
            gaps = pd.Series(changed_dates).diff().dt.total_seconds().div(86400).dropna()
            median = float(gaps.median()) if len(gaps) else np.nan
            comparisons = len(valid) - 1
            unchanged = int((~changes).sum())
            records.append(dict(
                Frequency=frequency, Asset=asset, Total_rows=len(s),
                First_valid_date=first, Last_valid_date=last,
                Valid_observations=len(valid), Missing_count_total=int(s.isna().sum()),
                Leading_missing_count=leading, Internal_missing_count=internal,
                Trailing_missing_count=trailing, Comparison_count=comparisons,
                Repeated_unchanged_count=unchanged,
                Repeated_unchanged_pct=100 * unchanged / comparisons if comparisons else np.nan,
                Changed_count=int(changes.sum()),
                Median_days_between_actual_updates=median,
                Mean_days_between_actual_updates=float(gaps.mean()) if len(gaps) else np.nan,
                Effective_updates_per_year=365.25 / gaps.mean() if len(gaps) else np.nan))
    return pd.DataFrame(records)


# Q1(b): construct complete common-sample monthly simple returns
def common_returns(prices):
    start = max(s.first_valid_index() for _, s in prices.items())
    end = min(s.last_valid_index() for _, s in prices.items())
    common = prices.loc[start:end].copy()
    if len(common) < 2 or common.isna().any().any():
        raise ValueError("Common price sample is too short or contains missing values; no filling performed")
    periods = common.index.to_period("M")
    expected = pd.period_range(periods[0], periods[-1], freq="M")
    if not periods.equals(expected):
        raise ValueError("Monthly dates are duplicated within a month or have skipped months")
    returns = common.pct_change(fill_method=None).iloc[1:]
    if not np.isfinite(returns.to_numpy()).all() or (returns <= -1).any().any():
        raise ValueError("Invalid simple returns; no rows silently removed")
    return common, returns


# Q1(c): sample moments (ddof=1); correlation is not annualised
def estimate(returns):
    if len(returns) < 2:
        raise ValueError("At least two return observations required")
    sd = returns.std(ddof=1)
    if (sd <= 0).any():
        raise ValueError("Zero-variance asset prevents a defined correlation matrix")
    params = pd.DataFrame({"Mean_Monthly_Return": returns.mean(),
                           "Annualised_Expected_Return": 12 * returns.mean(),
                           "Monthly_Volatility": sd,
                           "Annualised_Volatility": np.sqrt(12) * sd})
    params.index.name = "Asset"
    covariance = returns.cov(ddof=1)
    return params, covariance, 12 * covariance, returns.corr()


# Q1(d): diagnose without modifying C
def validate_covariance(covariance):
    c = covariance.to_numpy()
    asymmetry = float(np.max(np.abs(c - c.T)))
    scale = max(float(np.linalg.norm(c, ord=2)), np.finfo(float).tiny)
    tolerance = 100 * np.finfo(float).eps * len(c) * scale
    symmetric = asymmetry <= tolerance
    # eigvalsh is valid only for a symmetric matrix; fail explicitly otherwise.
    if not symmetric:
        raise ValueError(f"Covariance is asymmetric: {asymmetry}; tolerance {tolerance}")
    eigenvalues = np.linalg.eigvalsh(c)
    psd = bool(eigenvalues.min() >= -tolerance)
    table = pd.DataFrame({"Metric": ["Maximum_absolute_asymmetry", "Numerical_tolerance",
        "Symmetric", "Minimum_eigenvalue", "Positive_semidefinite", "Matrix_modified"],
        "Value": [asymmetry, tolerance, symmetric, float(eigenvalues.min()), psd, False]})
    return table, eigenvalues, psd


# Q1(e): recent estimation periods and rolling 60-month estimates
def stability(returns, baseline):
    periods = {"Full common sample": returns}
    end = returns.index[-1]
    for years in (10, 5):
        cutoff = end - pd.DateOffset(years=years)
        recent = returns.loc[returns.index > cutoff]
        if len(recent) == 12 * years:
            periods[f"Recent {years} years"] = recent
    results, comparisons, matrix_comparisons = {}, [], []
    upper = np.triu_indices(returns.shape[1], k=1)
    base_corr, base_cov = baseline[3], baseline[2]
    for label, sample in periods.items():
        p, monthly_cov, c, corr = estimate(sample)
        results[label] = (p, c, corr)
        t = p.reset_index()
        t.insert(0, "Period", label)
        t["First_return_date"], t["Last_return_date"] = sample.index[0], sample.index[-1]
        t["Return_observations"] = len(sample)
        t["Mean_change_vs_full"] = p["Annualised_Expected_Return"].to_numpy() - baseline[0]["Annualised_Expected_Return"].to_numpy()
        t["Vol_change_vs_full"] = p["Annualised_Volatility"].to_numpy() - baseline[0]["Annualised_Volatility"].to_numpy()
        comparisons.append(t)
        delta = (corr - base_corr).to_numpy()[upper]
        matrix_comparisons.append(dict(Period=label, Return_observations=len(sample),
            Average_pairwise_correlation=float(corr.to_numpy()[upper].mean()),
            Mean_abs_correlation_change=float(np.abs(delta).mean()),
            Max_abs_correlation_change=float(np.abs(delta).max()),
            Relative_covariance_Frobenius_change=float(np.linalg.norm(c - base_cov) / np.linalg.norm(base_cov))))
    rolling_mean = 12 * returns.rolling(60, min_periods=60).mean().iloc[59:]
    rolling_vol = np.sqrt(12) * returns.rolling(60, min_periods=60).std(ddof=1).iloc[59:]
    scalar = []
    for i in range(59, len(returns)):
        window = returns.iloc[i-59:i+1]
        scalar.append(dict(Date=returns.index[i], First_return_date=window.index[0],
            Return_observations=60,
            Average_pairwise_correlation=float(window.corr().to_numpy()[upper].mean())))
    scalar = pd.DataFrame(scalar)
    if scalar.empty:
        raise ValueError("Fewer than 60 monthly returns: requested rolling analysis unavailable")
    ranges = pd.DataFrame({"Mean_min": rolling_mean.min(), "Mean_max": rolling_mean.max(),
                          "Vol_min": rolling_vol.min(), "Vol_max": rolling_vol.max()})
    ranges["Mean_range"] = ranges.Mean_max - ranges.Mean_min
    ranges["Vol_range"] = ranges.Vol_max - ranges.Vol_min
    ranges.index.name = "Asset"
    return results, pd.concat(comparisons, ignore_index=True), pd.DataFrame(matrix_comparisons), rolling_mean, rolling_vol, scalar, ranges


def save_figures(out, corr, mean, vol, scalar, comparisons):
    fig, ax = plt.subplots(figsize=(10, 8), layout="constrained")
    im = ax.imshow(corr, vmin=-1, vmax=1)
    labels = corr.columns.str.replace(" Index", "", regex=False)
    ax.set_xticks(range(9), labels, rotation=45, ha="right")
    ax.set_yticks(range(9), labels)
    for i in range(9):
        for j in range(9):
            ax.text(j, i, f"{corr.iloc[i,j]:.2f}", ha="center", va="center",
                    color="white" if corr.iloc[i,j] < 0.25 else "black", fontsize=9)
    ax.set_title("Common-sample monthly return correlations")
    fig.colorbar(im, ax=ax, label="Correlation")
    fig.savefig(out / "Q1_fig_correlation_heatmap.png", dpi=200)
    plt.close(fig)
    for table, name, title in [(mean, "expected_returns", "Annualised expected return"),
                               (vol, "volatility", "Annualised volatility")]:
        fig, axes = plt.subplots(3, 3, figsize=(13, 9), sharex=True, layout="constrained")
        for ax, asset in zip(axes.flat, table.columns):
            ax.plot(table.index, table[asset])
            ax.set_title(asset, fontsize=10)
            ax.yaxis.set_major_formatter(PercentFormatter(1))
            ax.grid(alpha=.25)
        fig.suptitle(f"{title}: rolling 60-month windows (individual y scales)")
        fig.savefig(out / f"Q1_fig_rolling_{name}.png", dpi=200)
        plt.close(fig)
    fig, axes = plt.subplots(2, 1, figsize=(11, 8), layout="constrained")
    axes[0].plot(scalar.Date, scalar.Average_pairwise_correlation)
    axes[0].set(title="Average of 36 distinct pairwise correlations: rolling 60 months", ylabel="Correlation")
    for label, group in comparisons.groupby("Period", sort=False):
        axes[1].plot(group.Asset.str.replace(" Index", "", regex=False), group.Annualised_Expected_Return, marker="o", label=label)
    axes[1].set(title="Annualised expected returns by estimation period", ylabel="Expected return")
    axes[1].yaxis.set_major_formatter(PercentFormatter(1))
    axes[1].tick_params(axis="x", rotation=30)
    axes[1].legend()
    for ax in axes:
        ax.grid(alpha=.25)
    fig.savefig(out / "Q1_fig_parameter_stability.png", dpi=200)
    plt.close(fig)


def write_workbook(path, tables):
    # pandas/openpyxl keeps the explicitly requested single Python entry point portable.
    # Outputs contain computed static estimates, not editable Excel calculation models.
    with pd.ExcelWriter(path, engine="openpyxl", datetime_format="yyyy-mm-dd") as writer:
        for name, (table, description, kind) in tables.items():
            table.to_excel(writer, sheet_name=name, index=False, startrow=3)
            ws = writer.sheets[name]
            ws.cell(2, 1, description)
            ws.cell(2, 1).font = Font(name="Arial", size=14, bold=True)
            ws.sheet_view.showGridLines = False
            ws.freeze_panes = "C5" if name == "Q1a_Data_Audit" else "B5"
            ws.auto_filter.ref = f"A4:{get_column_letter(ws.max_column)}{ws.max_row}"
            for cell in ws[4]:
                cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="334155")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.row_dimensions[4].height = 42
            for row in ws.iter_rows(min_row=5):
                for cell in row:
                    cell.font = Font(name="Arial", size=10)
                    cell.alignment = Alignment(vertical="center")
                    if isinstance(cell.value, (pd.Timestamp,)) or cell.is_date:
                        cell.number_format = "yyyy-mm-dd"
                    elif isinstance(cell.value, float):
                        header = str(ws.cell(4, cell.column).value)
                        if kind == "rates" and cell.column > 1 and not any(x in header for x in ["observations", "date", "Period"]):
                            cell.number_format = "0.00%"
                        elif kind == "cov":
                            cell.number_format = "0.000000"
                        elif kind == "validation":
                            cell.number_format = "0.000000E+00"
                        else:
                            cell.number_format = "0.0000"
            for j, column in enumerate(table.columns, 1):
                width = max(17, min(32, len(str(column)) + 2))
                if table[column].dtype == object:
                    width = min(100, max(width, max((len(str(v)) for v in table[column]), default=0) + 2))
                ws.column_dimensions[get_column_letter(j)].width = width
            for row in ws.iter_rows(min_row=5):
                height = 18
                for cell in row:
                    if isinstance(cell.value, str):
                        width = ws.column_dimensions[cell.column_letter].width
                        if len(cell.value) > width - 2:
                            cell.alignment = Alignment(wrap_text=True, vertical="center")
                            height = max(height, 15 * (1 + int(len(cell.value) / (width - 2))))
                ws.row_dimensions[row[0].row].height = height


def main():
    here = Path(__file__).resolve().parent
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=here.parent / "Bloomber_data.xlsx")
    parser.add_argument("--output-dir", type=Path, default=here)
    args = parser.parse_args()
    out = args.output_dir.resolve()
    out.mkdir(parents=True, exist_ok=True)
    source_hash = hashlib.sha256(args.input.read_bytes()).hexdigest()
    print("Q1(a): parsing and auditing Daily, Weekly and Monthly sheets")
    data = {f: read_prices(args.input, f) for f in FREQUENCIES}
    assets = list(data["Monthly"].columns)
    if any(list(p.columns) != assets for p in data.values()):
        raise ValueError("Asset names/order differ across source sheets")
    audit = audit_prices(data)
    freq = audit[["Asset", "Frequency", "Valid_observations", "Leading_missing_count",
        "Internal_missing_count", "Repeated_unchanged_pct", "Median_days_between_actual_updates"]].copy()
    prior_path = here / "Q1a_data_audit.xlsx"
    reconciliation = []
    if prior_path.exists():
        prior = pd.read_excel(prior_path, sheet_name="Asset_Audit").set_index(["Frequency", "Asset"])
        current = audit.set_index(["Frequency", "Asset"])
        for col in ["Valid_observations", "Missing_count_total", "Leading_missing_count", "Internal_missing_count",
                    "Trailing_missing_count", "Repeated_unchanged_pct", "Median_days_between_actual_updates"]:
            delta = float((current[col] - prior[col]).abs().max())
            reconciliation.append({"Metric": col, "Maximum_absolute_difference": delta, "Matches": delta <= 1e-10})
    print("Q1(b)-(d): constructing monthly returns, estimating and validating moments")
    prices, returns = common_returns(data["Monthly"])
    baseline = estimate(returns)
    p, monthly_cov, c, corr = baseline
    validation, eigenvalues, psd = validate_covariance(c)
    print("Q1(e): estimating recent periods and rolling 60-month windows")
    results, comparisons, matrices, rolling_mean, rolling_vol, scalar, ranges = stability(returns, baseline)
    return_summary = pd.DataFrame({"Metric": ["Common_price_start", "Common_price_end", "Price_observations",
        "First_return_date", "Last_return_date", "Return_observations", "Assets", "Invalid_return_count", "Zero_return_count"],
        "Value": [prices.index[0], prices.index[-1], len(prices), returns.index[0], returns.index[-1], len(returns),
                  ", ".join(assets), 0, int((returns == 0).sum().sum())]})
    checklist = pd.DataFrame([
        ["Q1(a)", "History, missingness, repeats and effective frequency", "Q1a_Data_Audit; Q1a_Frequency_Comparison"],
        ["Q1(b)", "Monthly simple returns, common period and justification", "Q1b_Return_Summary; clean returns CSV; run summary"],
        ["Q1(c)", "Annualised mean and covariance, volatility and correlation", "Q1c sheets; baseline CSV; heatmap"],
        ["Q1(d)", "Symmetry, all eigenvalues, PSD and numerical tolerance", "Q1d_Covariance_Validation; Q1d_Eigenvalues"],
        ["Q1(e)", "Recent periods, rolling estimates and portfolio implications", "Q1e sheets; rolling figures; run summary"]],
        columns=["Requirement", "Addressed", "Evidence"])
    tables = {}
    def add(name, frame, description, kind="numeric"):
        tables[name] = (frame, description, kind)
    add("Q1a_Data_Audit", audit, "Q1(a): history and effective observation frequency")
    add("Q1a_Frequency_Comparison", freq, "Q1(a): compact frequency comparison (% repeats, calendar days)")
    if reconciliation:
        add("Q1a_Audit_Reconciliation", pd.DataFrame(reconciliation), "Agreement with preliminary audit")
    add("Q1b_Return_Summary", return_summary, "Q1(b): common price and return samples")
    add("Q1c_Annualised_Parameters", p.reset_index(), "Q1(c): monthly and annualised parameters (decimal rates)", "rates")
    for name, matrix, desc, kind in [("Q1c_Monthly_Covariance", monthly_cov, "Monthly sample covariance (decimal return squared)", "cov"),
        ("Q1c_Covariance", c, "Annualised covariance C = 12 x monthly covariance", "cov"),
        ("Q1c_Correlation", corr, "Monthly return correlation (dimensionless)", "numeric")]:
        add(name, matrix.rename_axis("Asset").reset_index(), desc, kind)
    add("Q1d_Covariance_Validation", validation, "Q1(d): covariance validation; matrix unchanged", "validation")
    add("Q1d_Eigenvalues", pd.DataFrame({"Eigenvalue_number": range(1, 10), "Eigenvalue": eigenvalues}), "All eigenvalues of annualised C (ascending)", "validation")
    add("Q1e_Period_Comparison", comparisons, "Q1(e): period parameters; changes are decimal-rate differences", "rates")
    add("Q1e_Matrix_Comparison", matrices, "Q1(e): compact correlation and covariance comparison")
    for label, (_, pc, pr) in results.items():
        if label.startswith("Recent"):
            years = label.split()[1]
            add(f"Q1e_{years}y_Covariance", pc.rename_axis("Asset").reset_index(), f"Recent {years}-year annualised covariance", "cov")
            add(f"Q1e_{years}y_Correlation", pr.rename_axis("Asset").reset_index(), f"Recent {years}-year correlation")
    add("Q1e_Rolling_Summary", ranges.reset_index(), "Q1(e): ranges of rolling 60-month parameters", "rates")
    add("Q1e_Rolling_Means", rolling_mean.reset_index(), "Rolling annualised means; window ending at Date", "rates")
    add("Q1e_Rolling_Volatility", rolling_vol.reset_index(), "Rolling annualised volatility; window ending at Date", "rates")
    add("Q1e_Rolling_Correlation", scalar, "Rolling average correlation of 36 distinct asset pairs")
    add("Q1_Requirements_Check", checklist, "Official Question 1(a)-(e) coverage")
    notes = [
        ("Method", "Monthly simple returns; one complete common sample; no filling or zero deletion."),
        ("Annualisation", "Mean x12; sample covariance x12; sample standard deviation x sqrt(12); correlation unchanged."),
        ("Sample convention", "Restrict prices first, then calculate returns; first price is not a return observation."),
        ("Recent windows", "Select return end dates strictly after end minus 10/5 calendar years; require 120/60 returns."),
        ("Audit convention", "Compare successive valid values; update gaps exclude initial value and measure time between actual changes."),
        ("Missingness", "Leading missing rows show unavailable earlier history; internal gaps lie between first and last valid values."),
        ("Units", "Returns/volatility are decimal rates; audit repeat percentage is in 0-100 units; covariance is decimal return squared."),
        ("Refresh", "Static Python-computed results; rerun Q1_full_analysis.py to refresh."),
        ("Source", str(args.input.resolve())),
        ("Source SHA256", source_hash),
        ("Review", "Supplied files do not establish currency comparability or price/total-return treatment for every index; confirm definitions before portfolio use."),
        ("Runtime", f"Python {platform.python_version()}, pandas {pd.__version__}, numpy {np.__version__}, matplotlib {matplotlib.__version__}")]
    readme = pd.DataFrame(notes + [(name, desc) for name, (_, desc, _) in tables.items()], columns=["Item", "Description"])
    tables = {"README": (readme, "FMAT3888 Q1: methods and workbook guide", "numeric"), **tables}
    returns.to_csv(out / "Q1_clean_monthly_returns.csv", date_format="%Y-%m-%d")
    p.reset_index().to_csv(out / "Q1_baseline_parameters.csv", index=False)
    write_workbook(out / "Q1_results.xlsx", tables)
    save_figures(out, corr, rolling_mean, rolling_vol, scalar, comparisons)
    hedge = audit[audit.Asset == "HEDGNAV Index"]
    summary = ["FMAT3888 Question 1 numerical run summary (not final report prose)",
        "Assets: " + ", ".join(assets),
        "Monthly baseline follows the agreed common-frequency design. HEDGNAV unchanged observations:"]
    summary += [f"  {r.Frequency}: {r.Repeated_unchanged_pct:.4f}%; median change gap {r.Median_days_between_actual_updates:.1f} calendar days." for r in hedge.itertuples()]
    summary += [f"Monthly repeats across nine assets: {audit.loc[audit.Frequency == 'Monthly', 'Repeated_unchanged_count'].sum()}; internal missing observations: {audit.loc[audit.Frequency == 'Monthly', 'Internal_missing_count'].sum()}.",
        "Higher-frequency dates therefore do not always carry new information. Leading missingness denotes shorter available history.",
        f"Common price period: {prices.index[0].date()} to {prices.index[-1].date()}; {len(prices)} prices.",
        f"Return period: {returns.index[0].date()} to {returns.index[-1].date()}; {len(returns)} monthly returns.",
        "Simple return R_t = P_t/P_(t-1) - 1. No missing-value filling or zero-return deletion.",
        "Annualisation: mean x12; covariance x12; volatility x sqrt(12). Sample covariance/SD use ddof=1. Correlation unchanged.",
        validation.to_string(index=False, float_format=lambda v: f"{v:.12g}"), "Eigenvalues: " + ", ".join(f"{v:.12g}" for v in eigenvalues),
        "Parameter stability: period estimates (decimal rates):", comparisons.to_string(index=False),
        "Matrix stability summaries:", matrices.to_string(index=False),
        f"Rolling windows: {len(rolling_mean)}, from {rolling_mean.index[0].date()} to {rolling_mean.index[-1].date()}, 60 observations each.",
        f"Average pairwise rolling correlation range: {scalar.Average_pairwise_correlation.min():.4f} to {scalar.Average_pairwise_correlation.max():.4f}.",
        "Rolling annualised parameter ranges (decimal rates):", ranges.to_string(),
        "Implication: the displayed period differences and rolling ranges show that estimated inputs vary over time. Subsequent portfolio construction should compare allocations under these alternative inputs; optimisation sensitivity is not measured in Q1.",
        "Review: annualised means are arithmetic estimates, not CAGR. Calendar month-end source dates sometimes fall on the last business day.",
        "Review: unchanged monthly levels do not independently prove freshness or absence of index smoothing.",
        notes[10][1],
        "No covariance repair performed. " + ("Symmetry and PSD checks pass." if psd else "PSD FAIL: investigate parsing and numerical conditioning; a complete-sample sample covariance is theoretically PSD."),
        "Recent windows omitted if insufficient observations: " + str([y for y in (10, 5) if f'Recent {y} years' not in results]),
        "Preliminary audit reconciliation: " + (str(all(r['Matches'] for r in reconciliation)) if reconciliation else "not available"),
        "Official requirements checked:", checklist.to_string(index=False),
        "Source SHA256: " + source_hash]
    (out / "Q1_run_summary.txt").write_text("\n".join(summary) + "\n", encoding="utf-8")
    if hashlib.sha256(args.input.read_bytes()).hexdigest() != source_hash:
        raise RuntimeError("Input workbook changed during analysis")
    print(f"Complete: {len(returns)} monthly returns; PSD={psd}; outputs in {out}")
    if not psd:
        raise RuntimeError("PSD validation failed; diagnostic outputs saved, C not repaired")


if __name__ == "__main__":
    main()
